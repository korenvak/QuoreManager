"""
Excel/XLSX Catalog Handler for Kitchen Quote Management System
Handles reading catalog data from XLSX files with category support
Version 2.0 - Supports new format with units, approval, and comments
"""

import pandas as pd
import logging
import openpyxl
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
import io
import os
import tempfile
import json
import pickle

# Column mappings for different formats
NEW_FORMAT_MAPPING = {
    'NAME': 'שם מוצר',
    'CATEGORY': 'קטגוריה', 
    'UNIT': 'יחידה',
    'COST': 'מחיר',
    'MANAGER APPROVED': 'דורש אישור',
    'COMMENT': 'תיאור'
}

OLD_FORMAT_MAPPING = {
    'הפריט': 'שם מוצר',
    'יחידת מידה': 'יחידה', 
    'מחיר יחידה': 'מחיר',
    'דורש אישור': 'דורש אישור',
    'הערות': 'תיאור'
}

class CatalogHandler:
    """Handles XLSX catalog file operations with support for new format"""
    
    _instance = None
    _initialized = False
    
    def __new__(cls):
        """Singleton pattern implementation"""
        if cls._instance is None:
            cls._instance = super(CatalogHandler, cls).__new__(cls)
        return cls._instance
    
    def __init__(self):
        """Initialize catalog handler (singleton)"""
        if self._initialized:
            return
            
        self.logger = logging.getLogger(__name__)
        self.catalog_data = []
        self.categories = []
        self.images = {}  # Store extracted images
        self.catalog_file_path = None
        self.format_detected = None  # 'new', 'old', or 'mixed'
        
        # Move cache to AppData for EXE persistence
        self.cache_dir = self._get_cache_directory()
        self.load_cached_catalog()
        
        self._initialized = True
    
    def _get_cache_directory(self) -> Path:
        """Get cache directory in AppData for EXE persistence"""
        try:
            import os
            appdata = os.environ.get('APPDATA')
            if appdata:
                cache_dir = Path(appdata) / "KitchenQuotes" / "cache"
            else:
                cache_dir = Path("config") / "cache"
            
            cache_dir.mkdir(parents=True, exist_ok=True)
            return cache_dir
        except Exception:
            # Fallback to local directory
            cache_dir = Path("config")
            cache_dir.mkdir(exist_ok=True)
            return cache_dir
    
    def load_cached_catalog(self):
        """Load cached catalog data if available"""
        try:
            cache_file = self.cache_dir / "catalog_cache.pkl"
            if cache_file.exists():
                with open(cache_file, 'rb') as f:
                    cached_data = pickle.load(f)
                    self.catalog_data = cached_data.get('items', [])
                    self.categories = cached_data.get('categories', [])
                    self.catalog_file_path = cached_data.get('file_path')
                    self.format_detected = cached_data.get('format_detected')
                    self.logger.info(f"Loaded {len(self.catalog_data)} items from cache")
        except Exception as e:
            self.logger.warning(f"Failed to load cached catalog: {e}")
    
    def save_cached_catalog(self):
        """Save catalog data to cache"""
        try:
            cache_file = self.cache_dir / "catalog_cache.pkl"
            
            cached_data = {
                'items': self.catalog_data,
                'categories': self.categories,
                'file_path': self.catalog_file_path,
                'format_detected': self.format_detected
            }
            
            with open(cache_file, 'wb') as f:
                pickle.dump(cached_data, f)
                
            self.logger.info(f"Saved {len(self.catalog_data)} items to cache")
        except Exception as e:
            self.logger.error(f"Failed to save catalog cache: {e}")
    
    def detect_sheet_format(self, df: pd.DataFrame) -> Tuple[str, int, Dict[str, int]]:
        """Detect the format of a sheet and return format type, header row, and column mapping"""
        
        # Check first 10 rows for headers
        for row_idx in range(min(10, len(df))):
            row_values = [str(cell_val).upper().strip() for cell_val in df.iloc[row_idx] if pd.notna(cell_val)]
            
            # Check for new English format
            new_format_cols = list(NEW_FORMAT_MAPPING.keys())
            found_new = [col for col in new_format_cols if col in row_values]
            
            if len(found_new) >= 4:  # Need at least 4 required columns
                self.logger.info(f"Detected NEW format in row {row_idx}: {found_new}")
                
                # Create column mapping
                col_mapping = {}
                for i, cell_val in enumerate(df.iloc[row_idx]):
                    if pd.notna(cell_val) and str(cell_val).upper().strip() in NEW_FORMAT_MAPPING:
                        col_mapping[str(cell_val).upper().strip()] = i
                
                return 'new', row_idx, col_mapping
            
            # Check for old Hebrew format
            old_format_cols = list(OLD_FORMAT_MAPPING.keys())
            found_old = [col for col in old_format_cols if any(old_col in str(row_val) for row_val in row_values for old_col in old_format_cols)]
            
            if len(found_old) >= 2:  # Need at least 2 required columns
                self.logger.info(f"Detected OLD format in row {row_idx}: {found_old}")
                
                # Create column mapping for old format
                col_mapping = {}
                for i, cell_val in enumerate(df.iloc[row_idx]):
                    if pd.notna(cell_val):
                        val_str = str(cell_val).strip()
                        for old_col in OLD_FORMAT_MAPPING.keys():
                            if old_col in val_str:
                                col_mapping[old_col] = i
                                break
                
                return 'old', row_idx, col_mapping
        
        return 'unknown', 0, {}
    
    def load_catalog(self, file_path: str) -> bool:
        """Load catalog from XLSX file with new format support"""
        try:
            file_path_obj = Path(file_path)
            if not file_path_obj.exists():
                self.logger.error(f"Catalog file not found: {file_path}")
                return False
            
            # Clear existing data
            self.catalog_data = []
            self.categories = []
            self.images = {}
            self.catalog_file_path = file_path
            self.format_detected = None
            
            # Load workbook for image extraction
            workbook = openpyxl.load_workbook(file_path)
            
            formats_found = []
            
            # Process each worksheet
            for sheet_name in workbook.sheetnames:
                self.logger.info(f"Processing sheet: {sheet_name}")
                
                worksheet = workbook[sheet_name]
                self.extract_images_from_sheet(worksheet, sheet_name)
                
                # Read sheet without header to detect format
                df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
                
                # Detect format
                format_type, header_row, col_mapping = self.detect_sheet_format(df)
                formats_found.append(format_type)
                
                if format_type != 'unknown':
                    items = self.process_sheet_data_v2(df, sheet_name, format_type, header_row, col_mapping)
                    if items:
                        self.catalog_data.extend(items)
                else:
                    self.logger.warning(f"Unknown format in sheet {sheet_name}, trying fallback")
                    # Try old processing method as fallback
                    items = self.process_sheet_data(df, sheet_name)
                    if items:
                        self.catalog_data.extend(items)
            
            workbook.close()
            
            # Determine overall format
            if 'new' in formats_found:
                self.format_detected = 'new' if formats_found.count('new') > formats_found.count('old') else 'mixed'
            elif 'old' in formats_found:
                self.format_detected = 'old'
            else:
                self.format_detected = 'unknown'
            
            # Save to cache
            self.save_cached_catalog()
            
            self.logger.info(f"Loaded {len(self.catalog_data)} items from catalog (format: {self.format_detected})")
            return True
            
        except Exception as e:
            self.logger.error(f"Failed to load catalog: {e}")
            return False
    
    def extract_images_from_sheet(self, worksheet, sheet_name: str):
        """Extract images from worksheet"""
        try:
            if not hasattr(worksheet, '_images'):
                return
                
            for image_obj in worksheet._images:
                # Get image data
                image_data = image_obj._data()
                
                # Convert to PIL Image
                pil_image = Image.open(io.BytesIO(image_data))
                
                # Get anchor position
                anchor = image_obj.anchor
                if hasattr(anchor, '_from') and anchor._from:
                    col = anchor._from.col
                    row = anchor._from.row
                    
                    # Store image with position key
                    image_key = f"{sheet_name}_{row}_{col}"
                    
                    # Save image to temporary file
                    temp_dir = Path(tempfile.gettempdir()) / "kitchen_quotes_images"
                    temp_dir.mkdir(exist_ok=True)
                    
                    image_path = temp_dir / f"{image_key}.png"
                    pil_image.save(image_path, "PNG")
                    
                    self.images[image_key] = str(image_path)
                    
        except Exception as e:
            self.logger.warning(f"Failed to extract images from {sheet_name}: {e}")
    
    def process_sheet_data_v2(self, df: pd.DataFrame, sheet_name: str, format_type: str, header_row: int, col_mapping: Dict[str, int]) -> List[Dict[str, Any]]:
        """Process sheet data using new format detection - focused on English columns only"""
        items = []
        current_category = "ללא קטגוריה"
        
        try:
            # For new format, focus only on English columns
            if format_type == 'new':
                name_col = 'NAME'
                category_col = 'CATEGORY'
                cost_col = 'COST'
                unit_col = 'UNIT'
                approval_col = 'MANAGER APPROVED'
                comment_col = 'COMMENT'
            else:
                # Old format fallback
                name_col = 'הפריט'
                category_col = None
                cost_col = 'מחיר יחידה'
                unit_col = 'יחידת מידה'
                approval_col = 'דורש אישור'
                comment_col = 'הערות'
            
            # Process rows after header
            for idx in range(header_row + 1, len(df)):
                row = df.iloc[idx]
                
                # Extract values based on column mapping
                item_data = {}
                
                for col_key, col_idx in col_mapping.items():
                    if col_idx < len(row):
                        value = row.iloc[col_idx]
                        if pd.notna(value) and str(value).strip():
                            item_data[col_key] = str(value).strip()
                
                # Skip completely empty rows
                if not item_data:
                    continue
                
                name = item_data.get(name_col, '').strip()
                cost_str = item_data.get(cost_col, '').strip()
                
                # Skip rows without names
                if not name:
                    continue
                
                # If we have category column, use it to update current category
                if category_col and category_col in item_data:
                    category = item_data.get(category_col, '').strip()
                    if category and category != current_category:
                        current_category = category
                        if category not in self.categories:
                            self.categories.append(category)
                
                # Check if this is a category row (has name but is clearly a category)
                # For new format, we rely on the CATEGORY column, not inference
                if format_type != 'new' and name and not cost_str:
                    # Only for old format - treat as category
                    current_category = name
                    if current_category not in self.categories:
                        self.categories.append(current_category)
                    continue
                
                # Process as item row (has name, may or may not have cost)
                if name:
                    # Parse cost - handle empty/missing costs
                    cost_val = 0.0  # Default for custom pricing
                    if cost_str:
                        try:
                            # Clean cost string
                            clean_cost = str(cost_str).replace('₪', '').replace(',', '').replace('$', '').strip()
                            if clean_cost:
                                cost_val = float(clean_cost)
                        except (ValueError, TypeError):
                            cost_val = 0.0  # Set to 0 for custom pricing
                    
                    # Parse units - handle empty/missing units
                    units = item_data.get(unit_col, '').strip() if unit_col else ''
                    
                    # Normalize units
                    if units:
                        if units.lower() in ['מ"א', 'מטר', 'meter', 'linear meter', 'm', 'מ']:
                            units = 'מ"א'
                        elif units.lower() in ['יח\'', 'יח', 'piece', 'unit', 'pieces', 'pcs']:
                            units = 'יח׳'
                        elif units.lower() in ['מ"ר', 'מר', 'sqm', 'square meter', 'm2']:
                            units = 'מ"ר'
                        # Keep other units as-is
                    else:
                        # Empty units - set to empty string (quantity will be 1, no unit display in PDF)
                        units = ''
                    
                    # Parse approval requirement
                    approval = item_data.get(approval_col, '').strip().lower() if approval_col else ''
                    requires_approval = approval in ['n', 'no', 'false', '0', 'לא']  # 'n' means requires approval
                    
                    # Parse comments
                    comments = item_data.get(comment_col, '').strip() if comment_col else ''
                    
                    # Create item - always create if we have a name
                    item = {
                        'שם מוצר': name,
                        'קטגוריה': current_category,
                        'גיליון': sheet_name,
                        'שורה': idx + 1,
                        'כמות': 1,  # Default quantity
                        'מחיר': cost_val,  # 0 for custom pricing
                        'יחידה': units,  # Empty string for unitless items
                        'דורש אישור': requires_approval,
                        'תיאור': comments,
                        'תמונה': None,
                        'קוד מוצר': '',
                        'פורמט': format_type,  # Track which format this item came from
                        'has_custom_pricing': cost_val == 0.0,  # Flag for items needing custom pricing
                        'is_unitless': not bool(units)  # Flag for items without units
                    }
                    
                    # Look for image
                    image_key = f"{sheet_name}_{idx}"
                    if image_key in self.images:
                        item['תמונה'] = self.images[image_key]
                    else:
                        # Try nearby rows for image
                        for offset in [-1, 0, 1]:
                            nearby_key = f"{sheet_name}_{idx + offset}"
                            if nearby_key in self.images:
                                item['תמונה'] = self.images[nearby_key]
                                break
                    
                    items.append(item)
                    self.logger.debug(f"Added item: {name}, Cost: {cost_val}, Units: '{units}', Custom: {cost_val == 0.0}")
                    
        except Exception as e:
            self.logger.error(f"Error processing sheet {sheet_name} with new format: {e}")
        
        return items
    
    def process_sheet_data(self, df: pd.DataFrame, sheet_name: str) -> List[Dict[str, Any]]:
        """Process sheet data to extract items with categories using Hebrew headers"""
        items = []
        current_category = "ללא קטגוריה"
        header_row = None
        header_map = {}
        try:
            # Find the header row (look for 'הפריט' or 'הפריט ' and 'מחיר יחידה')
            for idx, row in df.iterrows():
                row_list = [str(cell).strip() if pd.notna(cell) else '' for cell in row]
                normalized = [cell.replace(' ', '') for cell in row_list]
                if (any(hdr in ['הפריט'] for hdr in normalized) and
                    any(hdr in ['מחיריחידה'] for hdr in normalized)):
                    header_row = idx
                    for i, col in enumerate(row_list):
                        col_clean = col.strip().replace(' ', '')
                        if col_clean:
                            header_map[col_clean] = i
                    break
            if header_row is None:
                self.logger.error(f"No header row found in sheet {sheet_name}")
                return items
            def safe_int(val):
                try:
                    return int(val)
                except Exception:
                    return None
            # Process rows after header
            for idx, row in df.iterrows():
                idx_int = safe_int(idx)
                header_row_int = safe_int(header_row)
                if idx_int is None or header_row_int is None:
                    continue
                if idx_int <= header_row_int:
                    continue
                row_list = [str(cell).strip() if pd.notna(cell) else '' for cell in row]
                # Category row: only item column has value, no price
                item_col = header_map.get('הפריט', header_map.get('הפריט', header_map.get('הפריט', 1)))
                price_col = header_map.get('מחיריחידה', header_map.get('המחיר', 4))
                desc_col = header_map.get('הערות', 6)
                name = row_list[item_col] if item_col is not None and item_col < len(row_list) else ''
                price = row_list[price_col] if price_col is not None and price_col < len(row_list) else ''
                description = row_list[desc_col] if desc_col is not None and desc_col < len(row_list) else ''
                # Detect category row
                if name and not price:
                    current_category = name
                    if current_category not in self.categories:
                        self.categories.append(current_category)
                    continue
                if name and price:
                    try:
                        price_val = float(str(price).replace('₪', '').replace(',', '').replace('$', '').strip())
                    except Exception:
                        price_val = 0.0
                    item = {
                        'שם מוצר': name,
                        'קטגוריה': current_category,
                        'גיליון': sheet_name,
                        'שורה': idx_int + 1,
                        'כמות': 1,
                        'מחיר': price_val,
                        'תמונה': None,
                        'תיאור': description,
                        'קוד מוצר': '',
                        'יחידה': 'יח׳'
                    }
                    items.append(item)
        except Exception as e:
            self.logger.error(f"Error processing sheet {sheet_name}: {e}")
        return items
    
    def is_category_row(self, row_data: List[str]) -> bool:
        """Check if row represents a category"""
        if not row_data or not row_data[0]:
            return False
        
        # Category rows typically have text in first column but no numeric price
        first_cell = row_data[0].strip()
        
        # Skip if first cell is empty or looks like an item code
        if not first_cell or first_cell.lower() in ['', 'nan', 'none']:
            return False
        
        # Check if there's no price in typical price columns
        has_price = False
        for i in range(1, min(len(row_data), 5)):  # Check first few columns for price
            cell_value = row_data[i].strip()
            if cell_value and self.is_price_value(cell_value):
                has_price = True
                break
        
        # Category if has text but no price
        return not has_price and len(first_cell) > 1
    
    def is_item_row(self, row_data: List[str]) -> bool:
        """Check if row represents an item"""
        if not row_data or len(row_data) < 2:
            return False
        
        # Item rows should have name and price
        name = row_data[0].strip() if row_data[0] else ""
        
        # Skip if no name or name looks like category
        if not name or name.lower() in ['', 'nan', 'none']:
            return False
        
        # Look for price in the row
        has_price = False
        for cell_value in row_data[1:]:
            if cell_value and self.is_price_value(str(cell_value).strip()):
                has_price = True
                break
        
        return has_price
    
    def is_price_value(self, value: str) -> bool:
        """Check if value looks like a price"""
        if not value:
            return False
        
        # Remove common currency symbols and separators
        clean_value = value.replace('₪', '').replace('$', '').replace(',', '').strip()
        
        try:
            float(clean_value)
            return True
        except ValueError:
            return False
    
    def parse_item_row(self, row_data: List[str], category: str, sheet_name: str, row_index: int) -> Optional[Dict[str, Any]]:
        """Parse item row into structured data"""
        try:
            if len(row_data) < 2:
                return None
            
            item = {
                'שם מוצר': row_data[0].strip(),
                'קטגוריה': category,
                'גיליון': sheet_name,
                'שורה': row_index + 1,  # Excel row number (1-based)
                'כמות': 1,  # Default quantity
                'מחיר': 0.0,
                'תמונה': None,
                'תיאור': '',
                'קוד מוצר': '',
                'יחידה': 'יח׳'
            }
            
            # Find price column - look for first numeric value
            for i, cell_value in enumerate(row_data[1:], 1):
                if cell_value and self.is_price_value(str(cell_value)):
                    try:
                        price_str = str(cell_value).replace('₪', '').replace('$', '').replace(',', '').strip()
                        item['מחיר'] = float(price_str)
                        break
                    except ValueError:
                        continue
            
            # Look for additional data in other columns
            if len(row_data) > 2:
                # Try to find description, product code, etc.
                for i, cell_value in enumerate(row_data[2:], 2):
                    if cell_value and str(cell_value).strip() and not self.is_price_value(str(cell_value)):
                        if not item['תיאור']:
                            item['תיאור'] = str(cell_value).strip()
                        elif not item['קוד מוצר']:
                            item['קוד מוצר'] = str(cell_value).strip()
            
            # Look for image associated with this row
            image_key = f"{sheet_name}_{row_index}"
            if image_key in self.images:
                item['תמונה'] = self.images[image_key]
            else:
                # Try nearby rows for image
                for offset in [-1, 0, 1]:
                    nearby_key = f"{sheet_name}_{row_index + offset}"
                    if nearby_key in self.images:
                        item['תמונה'] = self.images[nearby_key]
                        break
            
            return item
            
        except Exception as e:
            self.logger.error(f"Error parsing item row: {e}")
            return None
    
    def get_catalog_items(self) -> List[Dict[str, Any]]:
        """Get all catalog items"""
        return self.catalog_data
    
    def get_categories(self) -> List[str]:
        """Get all categories"""
        return self.categories
    
    def get_items_by_category(self, category: str) -> List[Dict[str, Any]]:
        """Get items filtered by category"""
        return [item for item in self.catalog_data if item.get('קטגוריה') == category]
    
    def search_items(self, search_term: str) -> List[Dict[str, Any]]:
        """Search items by name or description"""
        if not search_term:
            return self.catalog_data
        
        search_term = search_term.lower()
        results = []
        
        for item in self.catalog_data:
            # Search in name, description, and category
            searchable_text = ' '.join([
                item.get('שם מוצר', ''),
                item.get('תיאור', ''),
                item.get('קטגוריה', ''),
                item.get('קוד מוצר', '')
            ]).lower()
            
            if search_term in searchable_text:
                results.append(item)
        
        return results
    
    def get_item_by_name(self, name: str) -> Optional[Dict[str, Any]]:
        """Get specific item by name"""
        for item in self.catalog_data:
            if item.get('שם מוצר') == name:
                return item
        return None
    
    def validate_catalog_structure(self, file_path: str) -> Tuple[bool, List[str]]:
        """Validate catalog file structure"""
        issues = []
        
        try:
            # Check if file exists
            if not Path(file_path).exists():
                issues.append("קובץ הקטלוג לא נמצא")
                return False, issues
            
            # Try to read file
            try:
                workbook = openpyxl.load_workbook(file_path)
                if not workbook.sheetnames:
                    issues.append("הקובץ לא מכיל גיליונות עבודה")
                    return False, issues
                
                # Check each sheet
                for sheet_name in workbook.sheetnames:
                    df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
                    
                    if df.empty:
                        issues.append(f"גיליון '{sheet_name}' ריק")
                        continue
                    
                    # Check for basic structure
                    has_items = False
                    for index, row in df.iterrows():
                        row_data = [str(cell) if pd.notna(cell) else "" for cell in row]
                        if self.is_item_row(row_data):
                            has_items = True
                            break
                    
                    if not has_items:
                        issues.append(f"גיליון '{sheet_name}' לא מכיל פריטים תקינים")
                
                workbook.close()
                
            except Exception as e:
                issues.append(f"שגיאה בקריאת הקובץ: {str(e)}")
                return False, issues
            
            # If we got here, basic structure is OK
            if not issues:
                return True, []
            else:
                return len(issues) < len(workbook.sheetnames), issues  # Partial success
                
        except Exception as e:
            issues.append(f"שגיאה כללית: {str(e)}")
            return False, issues
    
    def export_catalog_template(self, output_path: str) -> bool:
        """Export catalog template file"""
        try:
            # Create sample data
            data = {
                'A': ['קטגוריה: ארונות בסיס', '', 'ארון בסיס 60 ס"מ', 'ארון בסיס 80 ס"מ', '', 
                      'קטגוריה: ארונות עליון', '', 'ארון עליון 60 ס"מ', 'ארון עליון 80 ס"מ'],
                'B': ['', '', '1,200', '1,500', '', '', '', '800', '1,000'],
                'C': ['', '', 'ארון בסיס סטנדרטי', 'ארון בסיס רחב', '', '', '', 'ארון עליון סטנדרטי', 'ארון עליון רחב'],
                'D': ['', '', 'B60', 'B80', '', '', '', 'U60', 'U80']
            }
            
            df = pd.DataFrame(data)
            df.to_excel(output_path, index=False, header=False)
            
            self.logger.info(f"Catalog template exported to {output_path}")
            return True
            
        except Exception as e:
            self.logger.error(f"Failed to export template: {e}")
            return False
    
    def clear_temp_images(self):
        """Clear temporary image files"""
        try:
            temp_dir = Path(tempfile.gettempdir()) / "kitchen_quotes_images"
            if temp_dir.exists():
                for file_path in temp_dir.glob("*.png"):
                    try:
                        file_path.unlink()
                    except Exception:
                        pass
        except Exception:
            pass
    
    def get_catalog_stats(self) -> Dict[str, Any]:
        """Get catalog statistics"""
        if not self.catalog_data:
            return {
                'total_items': 0,
                'categories': 0,
                'items_with_images': 0,
                'average_price': 0.0,
                'price_range': (0.0, 0.0)
            }
        
        prices = [item.get('מחיר', 0) for item in self.catalog_data if item.get('מחיר', 0) > 0]
        items_with_images = len([item for item in self.catalog_data if item.get('תמונה')])
        
        return {
            'total_items': len(self.catalog_data),
            'categories': len(self.categories),
            'items_with_images': items_with_images,
            'average_price': sum(prices) / len(prices) if prices else 0.0,
            'price_range': (min(prices), max(prices)) if prices else (0.0, 0.0)
        } 